from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy.exc import ProgrammingError
from sqlalchemy import select, or_
from typing import List, Optional
from decimal import Decimal, InvalidOperation
import openpyxl
from openpyxl.styles import Font, PatternFill
from io import BytesIO
import uuid

from backend.app.core.database import SessionLocal
from backend.app.models.user import User, RoleTier
from backend.app.models.solution import Solution
from backend.app.models.catalog import CatalogItem, CatalogItemKind
from backend.app.models.setting import AppSetting
from backend.app.api.deps import get_current_user, require_permission, get_db, has_permission
from backend.app.schemas.catalog import (
    CatalogItemResponse,
    CatalogItemCreate,
    CatalogItemUpdate,
    CatalogSettingsResponse,
    CatalogSettingsUpdate,
)

router = APIRouter()

CATALOG_ENTRY_COST_VISIBILITY_KEY = "catalog_entry_cost_visibility"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def get_catalog_entry_cost_visibility(db: Session) -> bool:

    try:
        with db.begin_nested():
            setting = db.execute(
                select(AppSetting).where(
                    AppSetting.key == CATALOG_ENTRY_COST_VISIBILITY_KEY
                )
            ).scalar_one_or_none()
        if setting is None:
            return True
        return bool(setting.bool_value)
    except ProgrammingError:
        return True


def can_access_catalog_cost(db: Session, user: User) -> bool:
    if has_permission("cost_visibility", user):
        return True
    if user.role_tier == RoleTier.CATALOG_ENTRY:
        return get_catalog_entry_cost_visibility(db)
    return False


def mask_cost_price(
    item: CatalogItem,
    user: User,
    allow_cost: Optional[bool] = None,
    db: Optional[Session] = None,
) -> CatalogItemResponse:
    resp = CatalogItemResponse.model_validate(item)
    if allow_cost is None:
        allow_cost = (
            can_access_catalog_cost(db, user)
            if db is not None
            else has_permission("cost_visibility", user)
        )
    if not allow_cost:
        resp.cost_price = None
    return resp


@router.get("/settings", response_model=CatalogSettingsResponse)
def get_catalog_settings(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("catalog_edit")),
):
    return CatalogSettingsResponse(
        catalog_entry_cost_visibility=get_catalog_entry_cost_visibility(db)
    )


@router.put("/settings", response_model=CatalogSettingsResponse)
def update_catalog_settings(
    payload: CatalogSettingsUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("super_admin")),
):
    try:
        with db.begin_nested():
            setting = db.execute(
                select(AppSetting).where(
                    AppSetting.key == CATALOG_ENTRY_COST_VISIBILITY_KEY
                )
            ).scalar_one_or_none()
            if setting is None:
                setting = AppSetting(
                    key=CATALOG_ENTRY_COST_VISIBILITY_KEY,
                    bool_value=payload.catalog_entry_cost_visibility,
                    updated_by_user_id=current_user.id,
                )
                db.add(setting)
            else:
                setting.bool_value = payload.catalog_entry_cost_visibility
                setting.updated_by_user_id = current_user.id
        db.commit()
        db.refresh(setting)
        return CatalogSettingsResponse(
            catalog_entry_cost_visibility=bool(setting.bool_value)
        )
    except ProgrammingError:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Database migration required: app_settings table is not yet applied.",
        )



@router.get("/export")
def export_catalog(
    solution_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("catalog_edit")),
):
    solution = db.execute(
        select(Solution).where(Solution.id == solution_id)
    ).scalar_one_or_none()
    if not solution:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Solution not found"
        )

    allow_cost = can_access_catalog_cost(db, current_user)
    items = (
        db.execute(
            select(CatalogItem)
            .where(CatalogItem.solution_id == solution_id)
            .order_by(CatalogItem.id.asc())
        )
        .scalars()
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog Export"

    headers = [
        "Action",
        "ID",
        "Lookup Key",
        "Kind",
        "Part Number",
        "Description",
        "Brand",
        "Sell Price",
        "Cost Price",
    ]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for item in items:
        kind_val = item.kind.value if hasattr(item.kind, "value") else str(item.kind)
        sell_val = float(item.sell_price) if item.sell_price is not None else None
        cost_val = (
            (float(item.cost_price) if item.cost_price is not None else None)
            if allow_cost
            else None
        )
        ws.append(
            [
                "UPDATE",
                item.id,
                item.lookup_key,
                kind_val,
                item.part_number,
                item.description,
                item.brand,
                sell_val,
                cost_val,
            ]
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_code = (solution.code or f"sol_{solution.id}").replace(" ", "_")
    return StreamingResponse(
        buf,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="catalog_export_{safe_code}.xlsx"'
        },
    )


@router.get("/template")
def download_catalog_template(
    current_user: User = Depends(require_permission("catalog_edit")),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog Import"

    headers = [
        "Action",
        "ID",
        "Lookup Key",
        "Kind",
        "Part Number",
        "Description",
        "Brand",
        "Sell Price",
        "Cost Price",
    ]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    ws.append(
        [
            "NEW",
            None,
            None,
            "reference",
            "EXAMPLE-PART-01",
            "Example Reference Broadcast Product",
            "Benchmark",
            1000.00,
            750.00,
        ]
    )

    ws_help = wb.create_sheet(title="Instructions")
    ws_help.append(["Column / Rule", "Guidance"])
    ws_help.append(
        [
            "Action",
            "Use 'NEW' to create a new Reference catalog item. Use 'UPDATE' (or leave blank) to update an existing catalog item.",
        ]
    )
    ws_help.append(
        [
            "ID / Lookup Key",
            "Recommended for updating existing items. If both are provided, they must match the same item in the selected solution.",
        ]
    )
    ws_help.append(
        [
            "Part Number Matching",
            "When ID and Lookup Key are omitted on an UPDATE row, Part Number is matched only if it uniquely identifies a single item in the solution. Ambiguous duplicate Part Numbers (e.g. BDLKDVQD2, PS1050, R31050) are rejected.",
        ]
    )
    ws_help.append(
        [
            "RULE_DRIVEN Items",
            "Items with Kind='rule_driven' cannot be created or deleted via import, and their Part Number and Kind cannot be changed.",
        ]
    )
    ws_help.append(
        [
            "Pricing & NULL vs 0.00",
            "Leave Cost Price blank on NEW rows to store NULL (unpriced). Enter 0.00 for zero-priced items. All rows are validated atomically.",
        ]
    )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": 'attachment; filename="catalog_import_template.xlsx"'
        },
    )


@router.get("/", response_model=List[CatalogItemResponse])
def search_catalog(
    solution_id: int,
    q: Optional[str] = None,
    part_number: Optional[str] = None,
    description: Optional[str] = None,
    brand: Optional[str] = None,
    kind: Optional[CatalogItemKind] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("sales_user")),
):
    query = select(CatalogItem).where(CatalogItem.solution_id == solution_id)

    if q:
        query = query.where(
            or_(
                CatalogItem.part_number.ilike(f"%{q}%"),
                CatalogItem.description.ilike(f"%{q}%"),
                CatalogItem.brand.ilike(f"%{q}%"),
            )
        )

    if part_number:
        query = query.where(CatalogItem.part_number.ilike(f"%{part_number}%"))
    if description:
        query = query.where(CatalogItem.description.ilike(f"%{description}%"))
    if brand:
        query = query.where(CatalogItem.brand.ilike(f"%{brand}%"))
    if kind:
        query = query.where(CatalogItem.kind == kind)

    query = query.order_by(CatalogItem.id.asc())
    items = db.execute(query).scalars().all()
    allow_cost = can_access_catalog_cost(db, current_user)
    return [mask_cost_price(item, current_user, allow_cost=allow_cost) for item in items]


@router.post("/", response_model=CatalogItemResponse, status_code=status.HTTP_201_CREATED)
def create_catalog_item(
    item_in: CatalogItemCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("catalog_edit")),
):
    solution = db.execute(
        select(Solution).where(Solution.id == item_in.solution_id)
    ).scalar_one_or_none()
    if not solution:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Solution not found"
        )

    allow_cost = can_access_catalog_cost(db, current_user)
    if not allow_cost and (
        "cost_price" in item_in.model_fields_set and item_in.cost_price is not None
    ):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Cost price modification is forbidden when Catalog Entry cost visibility is disabled",
        )

    clean_desc = item_in.description.strip()
    if not clean_desc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Description cannot be empty",
        )

    dump_data = item_in.model_dump()
    dump_data["description"] = clean_desc
    dump_data["part_number"] = (
        item_in.part_number.strip()
        if item_in.part_number and item_in.part_number.strip()
        else None
    )
    dump_data["brand"] = (
        item_in.brand.strip() if item_in.brand and item_in.brand.strip() else None
    )
    if not allow_cost:
        dump_data["cost_price"] = None
    dump_data["lookup_key"] = f"cat_usr_{uuid.uuid4().hex}"

    item = CatalogItem(**dump_data)
    db.add(item)
    db.commit()
    db.refresh(item)
    return mask_cost_price(item, current_user, allow_cost=allow_cost)


@router.post("/import")
def import_catalog(
    solution_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("catalog_edit")),
):
    solution = db.execute(
        select(Solution).where(Solution.id == solution_id)
    ).scalar_one_or_none()
    if not solution:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Solution not found"
        )

    try:
        contents = file.file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid Excel file format"
        )

    headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
    header_map = {
        str(h).strip().lower(): idx
        for idx, h in enumerate(headers)
        if h is not None and str(h).strip() != ""
    }

    required_cols = ["description"]
    for col in required_cols:
        if col not in header_map:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required column: {col}",
            )

    allow_cost = can_access_catalog_cost(db, current_user)
    errors = []
    created_count = 0
    updated_count = 0

    def is_blank(val) -> bool:
        return val is None or (isinstance(val, str) and val.strip() == "")

    try:
        with db.begin_nested():
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                if not row or all(is_blank(cell) for cell in row):
                    continue

                row_errors = []

                def get_val(*col_names):
                    for col_name in col_names:
                        idx = header_map.get(col_name)
                        if idx is not None and idx < len(row):
                            val = row[idx]
                            if not is_blank(val):
                                return val
                    return None

                raw_action = get_val("action")
                raw_id = get_val("id")
                raw_lookup = get_val("lookup key", "lookup_key")
                raw_kind = get_val("kind")
                raw_part = get_val("part number", "part_number")
                raw_desc = get_val("description")
                raw_brand = get_val("brand")
                raw_sell = get_val("sell price", "sell_price")
                raw_cost = get_val("cost price", "cost_price")

                action_str = (
                    str(raw_action).strip().upper()
                    if not is_blank(raw_action)
                    else ""
                )
                if action_str not in ("", "UPDATE", "NEW"):
                    row_errors.append(
                        f"Invalid Action '{raw_action}'. Must be 'UPDATE', 'NEW', or blank."
                    )

                desc_str = ""
                if is_blank(raw_desc):
                    row_errors.append("Description is required")
                else:
                    desc_str = str(raw_desc).strip()

                def parse_decimal(val, field_name):
                    if is_blank(val):
                        return None
                    try:
                        dec = Decimal(str(val).strip())
                        if dec < 0:
                            row_errors.append(f"{field_name} must be non-negative")
                        return dec
                    except InvalidOperation:
                        row_errors.append(f"Invalid decimal format for {field_name}")
                        return None

                sell_price = parse_decimal(raw_sell, "Sell Price")
                cost_price = parse_decimal(raw_cost, "Cost Price")

                if not allow_cost and not is_blank(raw_cost):
                    row_errors.append(
                        "Cost Price modification is forbidden when Catalog Entry cost visibility is disabled"
                    )

                part_number = (
                    str(raw_part).strip() if not is_blank(raw_part) else None
                )
                brand = str(raw_brand).strip() if not is_blank(raw_brand) else None
                lookup_key_val = (
                    str(raw_lookup).strip() if not is_blank(raw_lookup) else None
                )

                kind_str = (
                    str(raw_kind).strip().lower() if not is_blank(raw_kind) else None
                )
                if kind_str is not None and kind_str not in (
                    CatalogItemKind.REFERENCE.value,
                    CatalogItemKind.RULE_DRIVEN.value,
                ):
                    row_errors.append(
                        f"Invalid Kind '{raw_kind}'. Must be 'reference' or 'rule_driven'."
                    )

                item_id = None
                if not is_blank(raw_id):
                    try:
                        if isinstance(raw_id, int) and not isinstance(raw_id, bool):
                            item_id = raw_id
                        elif isinstance(raw_id, float) and raw_id.is_integer():
                            item_id = int(raw_id)
                        else:
                            item_id = int(str(raw_id).strip())
                    except ValueError:
                        row_errors.append("ID must be an integer")

                if action_str == "NEW":
                    if item_id is not None or lookup_key_val is not None:
                        row_errors.append(
                            "Action 'NEW' cannot specify an existing ID or Lookup Key"
                        )
                    if kind_str == CatalogItemKind.RULE_DRIVEN.value:
                        row_errors.append(
                            "Cannot create RULE_DRIVEN items via Excel import; new items must be 'reference'"
                        )

                    if row_errors:
                        errors.append({"row": row_idx, "errors": row_errors})
                        continue

                    new_item = CatalogItem(
                        solution_id=solution_id,
                        lookup_key=f"cat_usr_{uuid.uuid4().hex}",
                        kind=CatalogItemKind.REFERENCE,
                        part_number=part_number,
                        description=desc_str,
                        brand=brand,
                        sell_price=sell_price,
                        cost_price=cost_price if allow_cost else None,
                    )
                    db.add(new_item)
                    created_count += 1
                else:
                    # UPDATE flow (Action is blank or 'UPDATE')
                    if row_errors:
                        errors.append({"row": row_idx, "errors": row_errors})
                        continue

                    target_item = None
                    if item_id is not None and lookup_key_val is not None:
                        item_by_id = db.execute(
                            select(CatalogItem).where(CatalogItem.id == item_id)
                        ).scalar_one_or_none()
                        item_by_key = db.execute(
                            select(CatalogItem).where(
                                CatalogItem.lookup_key == lookup_key_val
                            )
                        ).scalar_one_or_none()
                        if not item_by_id:
                            row_errors.append(f"CatalogItem with ID {item_id} not found")
                        if not item_by_key:
                            row_errors.append(
                                f"CatalogItem with Lookup Key '{lookup_key_val}' not found"
                            )
                        if (
                            item_by_id
                            and item_by_key
                            and item_by_id.id != item_by_key.id
                        ):
                            row_errors.append(
                                f"Conflicting ID ({item_id}) and Lookup Key ('{lookup_key_val}') refer to different catalog items"
                            )
                        elif item_by_id and item_by_key:
                            target_item = item_by_id
                    elif item_id is not None:
                        target_item = db.execute(
                            select(CatalogItem).where(CatalogItem.id == item_id)
                        ).scalar_one_or_none()
                        if not target_item:
                            row_errors.append(f"CatalogItem with ID {item_id} not found")
                    elif lookup_key_val is not None:
                        target_item = db.execute(
                            select(CatalogItem).where(
                                CatalogItem.lookup_key == lookup_key_val
                            )
                        ).scalar_one_or_none()
                        if not target_item:
                            row_errors.append(
                                f"CatalogItem with Lookup Key '{lookup_key_val}' not found"
                            )
                    else:
                        if part_number is not None:
                            matches = (
                                db.execute(
                                    select(CatalogItem).where(
                                        CatalogItem.solution_id == solution_id,
                                        CatalogItem.part_number == part_number,
                                    )
                                )
                                .scalars()
                                .all()
                            )
                            if len(matches) == 1:
                                target_item = matches[0]
                            elif len(matches) > 1:
                                row_errors.append(
                                    f"Ambiguous duplicate Part Number '{part_number}' matches {len(matches)} items in this solution. Provide ID or Lookup Key."
                                )
                            else:
                                row_errors.append(
                                    f"No existing catalog item found with Part Number '{part_number}' in this solution. Use Action 'NEW' to create a new product."
                                )
                        else:
                            row_errors.append(
                                "Missing identifier: provide ID, Lookup Key, or unique Part Number to update an existing item, or set Action to 'NEW' to create a new item."
                            )

                    if row_errors or target_item is None:
                        errors.append({"row": row_idx, "errors": row_errors})
                        continue

                    if target_item.solution_id != solution_id:
                        row_errors.append(
                            f"CatalogItem (ID {target_item.id}) belongs to a different solution"
                        )
                        errors.append({"row": row_idx, "errors": row_errors})
                        continue

                    if target_item.kind == CatalogItemKind.RULE_DRIVEN:
                        if (
                            kind_str is not None
                            and kind_str != CatalogItemKind.RULE_DRIVEN.value
                        ):
                            row_errors.append("Cannot change Kind of RULE_DRIVEN item")
                        if (
                            part_number is not None
                            and part_number != target_item.part_number
                        ):
                            row_errors.append(
                                "Cannot modify Part Number of RULE_DRIVEN item"
                            )
                    else:
                        if kind_str == CatalogItemKind.RULE_DRIVEN.value:
                            row_errors.append(
                                "Cannot change Kind of REFERENCE item to RULE_DRIVEN via Excel import"
                            )

                    if row_errors:
                        errors.append({"row": row_idx, "errors": row_errors})
                        continue

                    if target_item.kind != CatalogItemKind.RULE_DRIVEN:
                        target_item.part_number = part_number
                    target_item.description = desc_str
                    target_item.brand = brand
                    if not is_blank(raw_sell):
                        target_item.sell_price = sell_price
                    if allow_cost and not is_blank(raw_cost):
                        target_item.cost_price = cost_price
                    updated_count += 1

            if errors:
                raise ValueError("Validation failed")

        db.commit()
    except ValueError:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=errors)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Import failed"
        )

    return {
        "message": "Import successful",
        "created_count": created_count,
        "updated_count": updated_count,
    }


@router.get("/{item_id}", response_model=CatalogItemResponse)
def get_catalog_item(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("sales_user")),
):
    item = db.execute(
        select(CatalogItem).where(CatalogItem.id == item_id)
    ).scalar_one_or_none()
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Catalog Item not found"
        )
    allow_cost = can_access_catalog_cost(db, current_user)
    return mask_cost_price(item, current_user, allow_cost=allow_cost)


@router.put("/{item_id}", response_model=CatalogItemResponse)
def update_catalog_item(
    item_id: int,
    item_in: CatalogItemUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("catalog_edit")),
):
    item = db.execute(
        select(CatalogItem).where(CatalogItem.id == item_id)
    ).scalar_one_or_none()
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Catalog Item not found"
        )

    allow_cost = can_access_catalog_cost(db, current_user)
    update_data = item_in.model_dump(exclude_unset=True)

    if not allow_cost and "cost_price" in update_data:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Cost price modification is forbidden when Catalog Entry cost visibility is disabled",
        )

    if "part_number" in update_data:
        update_data["part_number"] = (
            update_data["part_number"].strip()
            if update_data["part_number"] and update_data["part_number"].strip()
            else None
        )
    if "brand" in update_data:
        update_data["brand"] = (
            update_data["brand"].strip()
            if update_data["brand"] and update_data["brand"].strip()
            else None
        )
    if "description" in update_data and update_data["description"] is not None:
        clean_desc = update_data["description"].strip()
        if not clean_desc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Description cannot be empty",
            )
        update_data["description"] = clean_desc

    if item.kind == CatalogItemKind.RULE_DRIVEN:
        if "kind" in update_data and update_data["kind"] != CatalogItemKind.RULE_DRIVEN:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot change kind of RULE_DRIVEN items",
            )
        if "part_number" in update_data and update_data["part_number"] != item.part_number:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot change part_number of RULE_DRIVEN items",
            )

    for key, value in update_data.items():
        setattr(item, key, value)

    db.commit()
    db.refresh(item)
    return mask_cost_price(item, current_user, allow_cost=allow_cost)


@router.delete("/{item_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_catalog_item(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("catalog_edit")),
):
    item = db.execute(
        select(CatalogItem).where(CatalogItem.id == item_id)
    ).scalar_one_or_none()
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Catalog Item not found"
        )

    if item.kind == CatalogItemKind.RULE_DRIVEN:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Implementation safety rule derived from calculation-engine dependency: Deletion of RULE_DRIVEN items is globally blocked.",
        )

    from backend.app.models.quote import QuoteLineItem

    is_referenced = db.execute(
        select(QuoteLineItem).where(QuoteLineItem.catalog_item_id == item_id)
    ).first()
    if is_referenced:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cannot delete Catalog Item referenced by existing quotes",
        )

    db.delete(item)
    db.commit()
