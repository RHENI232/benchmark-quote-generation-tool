import pytest
import io
import openpyxl
from decimal import Decimal
from sqlalchemy import select
from backend.app.models.user import User, RoleTier, AccountType
from backend.app.core.security import get_password_hash
from backend.app.models.catalog import CatalogItemKind


def get_auth_headers(client, email, password):
    resp = client.post("/api/auth/login", json={"email": email, "password": password})
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


@pytest.fixture
def setup_data(client, db):
    users_data = [
        ("admin@test.com", RoleTier.ADMIN),
        ("mgmt@test.com", RoleTier.MANAGEMENT),
        ("catalog@test.com", RoleTier.CATALOG_ENTRY),
        ("sales@test.com", RoleTier.SALES),
    ]
    for email, role in users_data:
        db.execute(User.__table__.delete().where(User.email == email))
        user = User(
            email=email,
            role_tier=role,
            account_type=AccountType.MANUAL,
            password_hash=get_password_hash("pass"),
            enabled=True,
        )
        db.add(user)
    db.commit()

    headers_admin = get_auth_headers(client, "admin@test.com", "pass")

    # Ensure default catalog_entry_cost_visibility is True at the start of each test
    client.put(
        "/api/catalog/settings",
        json={"catalog_entry_cost_visibility": True},
        headers=headers_admin,
    )

    resp = client.post(
        "/api/solutions",
        json={"name": "Catalog Test Sol", "default_margin_percent": 15.0},
        headers=headers_admin,
    )
    sol_id = resp.json()["id"]

    return {"sol_id": sol_id, "admin": headers_admin}


def test_catalog_settings_route_not_captured_by_item_id(client, setup_data):
    headers_admin = setup_data["admin"]
    headers_mgmt = get_auth_headers(client, "mgmt@test.com", "pass")
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")
    headers_sales = get_auth_headers(client, "sales@test.com", "pass")

    # 1. GET /api/catalog/settings reaches settings endpoint (never 422 int_parsing on item_id)
    r_get = client.get("/api/catalog/settings", headers=headers_cat)
    assert r_get.status_code == 200
    assert r_get.json() == {"catalog_entry_cost_visibility": True}

    # Sales is forbidden from GET /api/catalog/settings (403, not 422 int_parsing)
    r_sales_get = client.get("/api/catalog/settings", headers=headers_sales)
    assert r_sales_get.status_code == 403

    # 2. PUT /api/catalog/settings reaches settings update endpoint with super_admin authorization
    r_put_false = client.put(
        "/api/catalog/settings",
        json={"catalog_entry_cost_visibility": False},
        headers=headers_admin,
    )
    assert r_put_false.status_code == 200
    assert r_put_false.json() == {"catalog_entry_cost_visibility": False}

    # Non-admin roles receive 403 Forbidden (not 422 int_parsing)
    assert (
        client.put(
            "/api/catalog/settings",
            json={"catalog_entry_cost_visibility": True},
            headers=headers_mgmt,
        ).status_code
        == 403
    )
    assert (
        client.put(
            "/api/catalog/settings",
            json={"catalog_entry_cost_visibility": True},
            headers=headers_cat,
        ).status_code
        == 403
    )

    # Restore to True via Admin and verify via GET
    r_put_true = client.put(
        "/api/catalog/settings",
        json={"catalog_entry_cost_visibility": True},
        headers=headers_admin,
    )
    assert r_put_true.status_code == 200
    assert r_put_true.json() == {"catalog_entry_cost_visibility": True}


def test_cost_visibility_setting_and_enforcement(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_admin = setup_data["admin"]
    headers_mgmt = get_auth_headers(client, "mgmt@test.com", "pass")
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")
    headers_sales = get_auth_headers(client, "sales@test.com", "pass")

    # 1. Default policy: catalog_entry_cost_visibility is True
    r_set = client.get("/api/catalog/settings", headers=headers_cat)
    assert r_set.status_code == 200
    assert r_set.json()["catalog_entry_cost_visibility"] is True


    # Non-admin cannot toggle setting
    assert (
        client.put(
            "/api/catalog/settings",
            json={"catalog_entry_cost_visibility": False},
            headers=headers_cat,
        ).status_code
        == 403
    )
    assert (
        client.put(
            "/api/catalog/settings",
            json={"catalog_entry_cost_visibility": False},
            headers=headers_mgmt,
        ).status_code
        == 403
    )

    # 2. When enabled, Catalog Entry can create item with cost_price and view it on catalog endpoints
    resp_create = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "Test Cost Visibility",
            "sell_price": 100.0,
            "cost_price": 80.0,
        },
        headers=headers_cat,
    )
    assert resp_create.status_code == 201
    assert float(resp_create.json()["cost_price"]) == 80.0
    item_id = resp_create.json()["id"]

    # Sales NEVER sees cost_price on catalog endpoints
    resp_sales = client.get(f"/api/catalog/{item_id}", headers=headers_sales)
    assert resp_sales.status_code == 200
    assert resp_sales.json()["cost_price"] is None

    # Catalog Entry sees cost_price when setting is True
    resp_cat_get = client.get(f"/api/catalog/{item_id}", headers=headers_cat)
    assert float(resp_cat_get.json()["cost_price"]) == 80.0

    # 3. Admin disables catalog_entry_cost_visibility
    r_disable = client.put(
        "/api/catalog/settings",
        json={"catalog_entry_cost_visibility": False},
        headers=headers_admin,
    )
    assert r_disable.status_code == 200
    assert r_disable.json()["catalog_entry_cost_visibility"] is False

    # Catalog Entry now gets cost_price masked on read (detail and list)
    assert (
        client.get(f"/api/catalog/{item_id}", headers=headers_cat).json()["cost_price"]
        is None
    )
    list_cat = client.get(
        f"/api/catalog?solution_id={sol_id}", headers=headers_cat
    ).json()
    assert all(i["cost_price"] is None for i in list_cat)

    # Management and Admin still see cost_price even when Catalog Entry visibility is False
    assert (
        float(
            client.get(f"/api/catalog/{item_id}", headers=headers_mgmt).json()[
                "cost_price"
            ]
        )
        == 80.0
    )
    assert (
        float(
            client.get(f"/api/catalog/{item_id}", headers=headers_admin).json()[
                "cost_price"
            ]
        )
        == 80.0
    )

    # 4. When disabled, Catalog Entry is forbidden from supplying cost_price on POST or PUT
    r_forbid_post = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "Unauthorized Cost Create",
            "sell_price": 120.0,
            "cost_price": 90.0,
        },
        headers=headers_cat,
    )
    assert r_forbid_post.status_code == 403

    r_forbid_put = client.put(
        f"/api/catalog/{item_id}",
        json={"description": "Try Changing Cost", "cost_price": 50.0},
        headers=headers_cat,
    )
    assert r_forbid_put.status_code == 403

    # Catalog Entry CAN update non-cost fields without overwriting existing cost_price
    r_allowed_put = client.put(
        f"/api/catalog/{item_id}",
        json={"description": "Updated Description Without Cost", "sell_price": 110.0},
        headers=headers_cat,
    )
    assert r_allowed_put.status_code == 200
    assert r_allowed_put.json()["cost_price"] is None
    # Verify underlying cost_price is still 80.0 via Admin
    assert (
        float(
            client.get(f"/api/catalog/{item_id}", headers=headers_admin).json()[
                "cost_price"
            ]
        )
        == 80.0
    )


def test_missing_cost_price_preserved_as_null_and_zero_preserved(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")
    headers_admin = setup_data["admin"]

    # 1. Create item with omitted cost_price -> MUST remain NULL (None), never silently calculated from margin
    resp_null = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "Unpriced Cost Item",
            "sell_price": 100.0,
        },
        headers=headers_cat,
    )
    assert resp_null.status_code == 201
    null_item_id = resp_null.json()["id"]

    resp_admin_null = client.get(f"/api/catalog/{null_item_id}", headers=headers_admin)
    assert resp_admin_null.json()["cost_price"] is None

    # 2. Create item with explicit 0.00 cost_price and 0.00 sell_price -> MUST preserve 0.00 (not NULL)
    resp_zero = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "Zero Priced Item",
            "sell_price": 0.0,
            "cost_price": 0.0,
        },
        headers=headers_cat,
    )
    assert resp_zero.status_code == 201
    zero_item_id = resp_zero.json()["id"]

    resp_admin_zero = client.get(f"/api/catalog/{zero_item_id}", headers=headers_admin)
    assert resp_admin_zero.json()["sell_price"] is not None
    assert float(resp_admin_zero.json()["sell_price"]) == 0.0
    assert resp_admin_zero.json()["cost_price"] is not None
    assert float(resp_admin_zero.json()["cost_price"]) == 0.0


def test_duplicate_and_null_part_numbers_preserved(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")

    # Two items with the same part_number (e.g. PS1050) and one with NULL part_number
    r1 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": "PS1050",
            "description": "Professional Services Variant A",
            "sell_price": 200.0,
        },
        headers=headers_cat,
    )
    r2 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": "PS1050",
            "description": "Professional Services Variant B",
            "sell_price": 300.0,
        },
        headers=headers_cat,
    )
    r3 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": None,
            "description": "Service Without Part Number",
            "sell_price": 150.0,
        },
        headers=headers_cat,
    )
    assert r1.status_code == 201
    assert r2.status_code == 201
    assert r3.status_code == 201
    assert r1.json()["id"] != r2.json()["id"]
    assert r1.json()["lookup_key"] != r2.json()["lookup_key"]
    assert r3.json()["part_number"] is None


def test_catalog_search(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")

    client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "UniqueMonitor 4K",
            "brand": "Sony",
            "part_number": "SN-4K-01",
        },
        headers=headers_cat,
    )

    client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.RULE_DRIVEN.value,
            "description": "Audio Mixer Unique",
            "brand": "Yamaha",
            "part_number": "YM-01",
        },
        headers=setup_data["admin"],
    )

    headers_sales = get_auth_headers(client, "sales@test.com", "pass")
    r1 = client.get(f"/api/catalog?solution_id={sol_id}&q=Unique", headers=headers_sales)
    assert len(r1.json()) == 2

    r2 = client.get(
        f"/api/catalog?solution_id={sol_id}&brand=Sony", headers=headers_sales
    )
    assert len(r2.json()) == 1
    assert r2.json()[0]["description"] == "UniqueMonitor 4K"

    r3 = client.get(
        f"/api/catalog?solution_id={sol_id}&part_number=YM", headers=headers_sales
    )
    assert len(r3.json()) == 1
    assert r3.json()[0]["part_number"] == "YM-01"

    r4 = client.get(
        f"/api/catalog?solution_id={sol_id}&kind=rule_driven", headers=headers_sales
    )
    assert len(r4.json()) == 1
    assert r4.json()[0]["part_number"] == "YM-01"


def test_rule_driven_protection(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_admin = setup_data["admin"]

    resp = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.RULE_DRIVEN.value,
            "description": "Engine Core",
            "part_number": "ENG-001",
        },
        headers=headers_admin,
    )
    item_id = resp.json()["id"]

    r_del = client.delete(f"/api/catalog/{item_id}", headers=headers_admin)
    assert r_del.status_code == 400
    assert "Implementation safety rule" in r_del.json()["detail"]

    r_update = client.put(
        f"/api/catalog/{item_id}",
        json={"description": "Updated", "part_number": "ENG-002"},
        headers=headers_admin,
    )
    assert r_update.status_code == 400

    r_update_kind = client.put(
        f"/api/catalog/{item_id}",
        json={"description": "Updated", "kind": "reference"},
        headers=headers_admin,
    )
    assert r_update_kind.status_code == 400

    r_update2 = client.put(
        f"/api/catalog/{item_id}",
        json={"description": "Updated", "part_number": "ENG-001"},
        headers=headers_admin,
    )
    assert r_update2.status_code == 200


def test_export_and_template_endpoints(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_admin = setup_data["admin"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")
    headers_sales = get_auth_headers(client, "sales@test.com", "pass")

    client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": "EXP-01",
            "description": "Exportable Item",
            "brand": "Benchmark",
            "sell_price": 500.0,
            "cost_price": 350.0,
        },
        headers=headers_admin,
    )

    # Sales is forbidden from export and template
    assert (
        client.get(
            f"/api/catalog/export?solution_id={sol_id}", headers=headers_sales
        ).status_code
        == 403
    )
    assert client.get("/api/catalog/template", headers=headers_sales).status_code == 403

    # Template download works for Catalog Entry
    r_tpl = client.get("/api/catalog/template", headers=headers_cat)
    assert r_tpl.status_code == 200
    wb_tpl = openpyxl.load_workbook(io.BytesIO(r_tpl.content))
    assert "Catalog Import" in wb_tpl.sheetnames
    assert "Instructions" in wb_tpl.sheetnames

    # Export with cost visibility enabled includes Cost Price
    r_exp_en = client.get(
        f"/api/catalog/export?solution_id={sol_id}", headers=headers_cat
    )
    assert r_exp_en.status_code == 200
    wb_en = openpyxl.load_workbook(io.BytesIO(r_exp_en.content))
    ws_en = wb_en.active
    rows_en = list(ws_en.iter_rows(min_row=2, values_only=True))
    assert len(rows_en) == 1
    assert rows_en[0][4] == "EXP-01"
    assert float(rows_en[0][7]) == 500.0
    assert float(rows_en[0][8]) == 350.0

    # Disable cost visibility for Catalog Entry -> Export masks Cost Price
    client.put(
        "/api/catalog/settings",
        json={"catalog_entry_cost_visibility": False},
        headers=headers_admin,
    )
    r_exp_dis = client.get(
        f"/api/catalog/export?solution_id={sol_id}", headers=headers_cat
    )
    assert r_exp_dis.status_code == 200
    wb_dis = openpyxl.load_workbook(io.BytesIO(r_exp_dis.content))
    ws_dis = wb_dis.active
    rows_dis = list(ws_dis.iter_rows(min_row=2, values_only=True))
    assert rows_dis[0][8] is None


def test_import_catalog_new_and_update_matching(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")
    headers_admin = setup_data["admin"]

    # Seed an existing item to test UPDATE by ID, Lookup Key, and unique Part Number
    r_ex1 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": "UNIQ-PN-1",
            "description": "Original 1",
            "sell_price": 100.0,
            "cost_price": 70.0,
        },
        headers=headers_cat,
    ).json()
    r_ex2 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": "UNIQ-PN-2",
            "description": "Original 2",
            "sell_price": 200.0,
            "cost_price": 140.0,
        },
        headers=headers_cat,
    ).json()
    r_ex3 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "part_number": "UNIQ-PN-3",
            "description": "Original 3",
            "sell_price": 300.0,
            "cost_price": 210.0,
        },
        headers=headers_cat,
    ).json()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "Action",
            "ID",
            "Lookup Key",
            "Kind",
            "Part Number",
            "Description",
            "Brand",
            "Sell Price",
            "Cost Price",
        ]
    )
    # 1. Explicit NEW with cost_price supplied
    ws.append(
        ["NEW", "", "", "reference", "NEW-01", "Imported Ref", "TestBrand", 500, 400]
    )
    # 2. Explicit NEW with blank cost_price -> must preserve NULL (None)
    ws.append(
        ["NEW", "", "", "reference", "NEW-NULL-COST", "Unpriced Cost Ref", "BrandB", 250, ""]
    )
    # 3. UPDATE by ID
    ws.append(
        ["UPDATE", r_ex1["id"], "", "reference", "UNIQ-PN-1", "Updated 1", "B1", 110, 75]
    )
    # 4. UPDATE by Lookup Key
    ws.append(
        [
            "UPDATE",
            "",
            r_ex2["lookup_key"],
            "reference",
            "UNIQ-PN-2",
            "Updated 2",
            "B2",
            220,
            150,
        ]
    )
    # 5. UPDATE by unique Part Number (when ID and Lookup Key are omitted)
    ws.append(["UPDATE", "", "", "reference", "UNIQ-PN-3", "Updated 3", "B3", 330, 225])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        f"/api/catalog/import?solution_id={sol_id}",
        files={
            "file": (
                "import.xlsx",
                out,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers_cat,
    )
    assert resp.status_code == 200
    assert resp.json()["created_count"] == 2
    assert resp.json()["updated_count"] == 3

    # Verify NEW-01
    items_new1 = client.get(
        f"/api/catalog?solution_id={sol_id}&part_number=NEW-01", headers=headers_admin
    ).json()
    assert len(items_new1) == 1
    assert items_new1[0]["description"] == "Imported Ref"
    assert float(items_new1[0]["cost_price"]) == 400.0

    # Verify NEW-NULL-COST kept cost_price as None (not calculated from default margin!)
    items_null = client.get(
        f"/api/catalog?solution_id={sol_id}&part_number=NEW-NULL-COST",
        headers=headers_admin,
    ).json()
    assert len(items_null) == 1
    assert items_null[0]["cost_price"] is None

    # Verify updates
    assert (
        client.get(f"/api/catalog/{r_ex1['id']}", headers=headers_admin).json()[
            "description"
        ]
        == "Updated 1"
    )
    assert (
        client.get(f"/api/catalog/{r_ex2['id']}", headers=headers_admin).json()[
            "description"
        ]
        == "Updated 2"
    )
    assert (
        client.get(f"/api/catalog/{r_ex3['id']}", headers=headers_admin).json()[
            "description"
        ]
        == "Updated 3"
    )


def test_import_catalog_strict_validation_and_rollback(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_admin = setup_data["admin"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")

    # Create another solution to test cross-solution rejection
    other_sol_id = client.post(
        "/api/solutions",
        json={"name": "Other Solution For Import", "default_margin_percent": 15.0},
        headers=headers_admin,
    ).json()["id"]
    other_item = client.post(
        "/api/catalog",
        json={
            "solution_id": other_sol_id,
            "kind": "reference",
            "part_number": "OTHER-PN",
            "description": "Other Sol Item",
        },
        headers=headers_admin,
    ).json()

    # Create two items with duplicate Part Number 'PS1050' in sol_id
    dup1 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": "reference",
            "part_number": "PS1050",
            "description": "PS1050 Variant 1",
        },
        headers=headers_admin,
    ).json()
    dup2 = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": "reference",
            "part_number": "PS1050",
            "description": "PS1050 Variant 2",
        },
        headers=headers_admin,
    ).json()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "Action",
            "ID",
            "Lookup Key",
            "Kind",
            "Part Number",
            "Description",
            "Brand",
            "Sell Price",
            "Cost Price",
        ]
    )
    # Row 2: Valid NEW row (must be rolled back because subsequent rows fail!)
    ws.append(["NEW", "", "", "reference", "ROLLBACK-ME", "Valid Row", "B", 100, 80])
    # Row 3: Conflicting ID and Lookup Key
    ws.append(
        [
            "UPDATE",
            dup1["id"],
            dup2["lookup_key"],
            "reference",
            "PS1050",
            "Conflict Row",
            "B",
            100,
            80,
        ]
    )
    # Row 4: Cross-solution ID reference
    ws.append(
        [
            "UPDATE",
            other_item["id"],
            "",
            "reference",
            "OTHER-PN",
            "Cross Sol Row",
            "B",
            100,
            80,
        ]
    )
    # Row 5: Unknown ID
    ws.append(
        ["UPDATE", 999999, "", "reference", "UNK-PN", "Unknown ID Row", "B", 100, 80]
    )
    # Row 6: Ambiguous duplicate Part Number without ID/Lookup Key
    ws.append(
        ["UPDATE", "", "", "reference", "PS1050", "Ambiguous PN Row", "B", 100, 80]
    )
    # Row 7: Unmatched Part Number without Action='NEW' (must not silently create!)
    ws.append(
        ["UPDATE", "", "", "reference", "DOES-NOT-EXIST", "Silent Create Attempt", "B", 100, 80]
    )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        f"/api/catalog/import?solution_id={sol_id}",
        files={
            "file": (
                "import.xlsx",
                out,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers_cat,
    )
    assert resp.status_code == 400
    errors = resp.json()["detail"]
    error_rows = [e["row"] for e in errors]
    assert error_rows == [3, 4, 5, 6, 7]

    # Verify Row 2 ('ROLLBACK-ME') was rolled back completely
    r_check = client.get(
        f"/api/catalog?solution_id={sol_id}&part_number=ROLLBACK-ME",
        headers=headers_cat,
    )
    assert len(r_check.json()) == 0


def test_import_forbids_cost_price_when_visibility_disabled(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_admin = setup_data["admin"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")

    client.put(
        "/api/catalog/settings",
        json={"catalog_entry_cost_visibility": False},
        headers=headers_admin,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Action", "Part Number", "Description", "Sell Price", "Cost Price"])
    ws.append(["NEW", "FORBID-COST-1", "Cannot Set Cost", 100, 75])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    resp = client.post(
        f"/api/catalog/import?solution_id={sol_id}",
        files={
            "file": (
                "import.xlsx",
                out,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers_cat,
    )
    assert resp.status_code == 400
    errors = resp.json()["detail"]
    assert len(errors) == 1
    assert any("Cost Price modification is forbidden" in msg for msg in errors[0]["errors"])


def test_delete_referenced_catalog_item_and_snapshot_isolation(client, setup_data, db):
    sol_id = setup_data["sol_id"]
    headers_admin = setup_data["admin"]

    resp = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "Ref Item",
            "part_number": "REF-001",
            "sell_price": 500.0,
            "cost_price": 300.0,
        },
        headers=headers_admin,
    )
    item_id = resp.json()["id"]

    from backend.app.models.quote import Quote, QuoteLineItem
    from backend.app.models.region import Region

    region = Region(country_name="Test Region", currency_code="USD")
    db.add(region)
    db.commit()
    db.refresh(region)

    user = db.execute(select(User).where(User.email == "admin@test.com")).scalar_one()

    quote = Quote(
        region_id=region.id,
        solution_id=sol_id,
        client_name="Test Client",
        quote_ref_no="Q-12345",
        created_by_user_id=user.id,
        last_edited_by_user_id=user.id,
    )
    db.add(quote)
    db.commit()
    db.refresh(quote)

    line_item = QuoteLineItem(
        quote_id=quote.id,
        catalog_item_id=item_id,
        description="Ref Item",
        unit_cost_price_snapshot=Decimal("300.00"),
        unit_sell_price_snapshot=Decimal("500.00"),
        quantity=1,
        section="Hardware",
    )
    db.add(line_item)
    db.commit()
    db.refresh(line_item)

    # 1. Deleting referenced item is blocked
    r_del = client.delete(f"/api/catalog/{item_id}", headers=headers_admin)
    assert r_del.status_code == 400
    assert "Cannot delete" in r_del.json()["detail"]

    # 2. Updating catalog prices does NOT mutate existing QuoteLineItem snapshots
    r_upd = client.put(
        f"/api/catalog/{item_id}",
        json={"sell_price": 999.0, "cost_price": 777.0},
        headers=headers_admin,
    )
    assert r_upd.status_code == 200
    db.refresh(line_item)
    assert line_item.unit_sell_price_snapshot == Decimal("500.00")
    assert line_item.unit_cost_price_snapshot == Decimal("300.00")


def test_catalog_identity_create_update(client, setup_data):
    sol_id = setup_data["sol_id"]
    headers_cat = get_auth_headers(client, "catalog@test.com", "pass")

    resp_create = client.post(
        "/api/catalog",
        json={
            "solution_id": sol_id,
            "kind": CatalogItemKind.REFERENCE.value,
            "description": "Test Identity",
            "lookup_key": "hacked_key",
        },
        headers=headers_cat,
    )
    assert resp_create.status_code == 201

    data = resp_create.json()
    assert "lookup_key" in data
    assert data["lookup_key"] != "hacked_key"
    assert data["lookup_key"].startswith("cat_usr_")

    item_id = data["id"]
    original_key = data["lookup_key"]

    resp_update = client.put(
        f"/api/catalog/{item_id}",
        json={
            "description": "Updated Identity",
            "lookup_key": "another_hacked_key",
        },
        headers=headers_cat,
    )

    assert resp_update.status_code == 200
    data_up = resp_update.json()
    assert data_up["description"] == "Updated Identity"
    assert data_up["lookup_key"] == original_key
